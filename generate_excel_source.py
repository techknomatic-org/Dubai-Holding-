import openpyxl
import json
from datetime import datetime

wb = openpyxl.load_workbook('DH_TechM_Managed_Services_BI_Source.xlsx', data_only=True)

def serialize_val(val):
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    if val is None:
        return None
    return val

# 1. Parse MOM Actions from 02_MOM_Actions
mom_sheet = wb['02_MOM_Actions']
mom_rows = list(mom_sheet.iter_rows(values_only=True))

# Header row is where row[0] == 'Action_No'
header_idx = -1
for i, r in enumerate(mom_rows):
    if r[0] == 'Action_No':
        header_idx = i
        break

headers = mom_rows[header_idx]
raw_actions = []
for r in mom_rows[header_idx+1:]:
    if r[0] is not None and str(r[0]).strip() != '':
        item = {h: serialize_val(v) for h, v in zip(headers, r) if h}
        raw_actions.append(item)

# 2. Parse Qualys from 21_Qualys_Project
qualys_sheet = wb['21_Qualys_Project']
qualys_rows = list(qualys_sheet.iter_rows(values_only=True))
# Summary row
qualys_summary = {}
for i, r in enumerate(qualys_rows):
    if r[0] == 'Reporting_Month':
        vals = qualys_rows[i+1]
        qualys_summary = {
            'reportingMonth': serialize_val(vals[0]),
            'completionPct': vals[1],
            'completedTaskCount': vals[2],
            'ongoingTaskCount': vals[3],
            'planStartDate': serialize_val(vals[4]),
            'planEndDate': serialize_val(vals[5])
        }
        break

# 3. Parse ITOps Monthly from 03_ITOps_Monthly
itops_sheet = wb['03_ITOps_Monthly']
itops_rows = list(itops_sheet.iter_rows(values_only=True))
itops_data = []
for i, r in enumerate(itops_rows):
    if r[0] == 'Month':
        itops_headers = r
        for row in itops_rows[i+1:]:
            if row[0] is not None:
                itops_data.append({h: serialize_val(v) for h, v in zip(itops_headers, row) if h})
        break

# 4. Parse Automation Monthly from 13_Automation_Monthly
auto_sheet = wb['13_Automation_Monthly']
auto_rows = list(auto_sheet.iter_rows(values_only=True))
auto_monthly = []
for i, r in enumerate(auto_rows):
    if r[0] == 'Month':
        auto_headers = r
        for row in auto_rows[i+1:]:
            if row[0] is not None:
                auto_monthly.append({h: serialize_val(v) for h, v in zip(auto_headers, row) if h})
        break

# 5. Parse Risk Dashboard Summary from 22_Risk_Dashboard
risk_sheet = wb['22_Risk_Dashboard']
risk_rows = list(risk_sheet.iter_rows(values_only=True))
risk_summary = {}
for i, r in enumerate(risk_rows):
    if r[0] == 'Metric' and r[1] == 'Value':
        for row in risk_rows[i+1:i+8]:
            if row[0]:
                risk_summary[str(row[0])] = row[1]
        break

# 6. Parse Value Savings from 24_Value_Savings
val_sheet = wb['24_Value_Savings']
val_rows = list(val_sheet.iter_rows(values_only=True))
savings_data = []
for i, r in enumerate(val_rows):
    if r[0] == 'Month':
        val_headers = r
        for row in val_rows[i+1:]:
            if row[0] is not None:
                savings_data.append({h: serialize_val(v) for h, v in zip(val_headers, row) if h})
        break

print("Extracted Excel components successfully:")
print("Actions count:", len(raw_actions))
print("Qualys summary:", qualys_summary)
print("ITOps months:", len(itops_data))
print("Risk summary:", risk_summary)

# Generate src/data/excelDataSource.ts
ts_content = f"""// AUTO-GENERATED DIRECTLY FROM DH_TechM_Managed_Services_BI_Source.xlsx
// 100% Sourced from Excel Sheets

export interface RawExcelMomAction {{
  Action_No: number;
  Action: string;
  Type: 'I' | 'A';
  Type_Description: string;
  Owner: string;
  Original_Due_Date: string;
  Revised_Due_Date: string | null;
  Status: 'Closed' | 'On going' | 'Open';
  Remarks: string | null;
  Source_Page: number;
}}

export interface ExcelQualysSummary {{
  reportingMonth: string;
  completionPct: number;
  completedTaskCount: number;
  ongoingTaskCount: number;
  planStartDate: string;
  planEndDate: string;
}}

export const EXCEL_MOM_ACTIONS: RawExcelMomAction[] = {json.dumps(raw_actions, indent=2)};

export const EXCEL_QUALYS_SUMMARY: ExcelQualysSummary = {json.dumps(qualys_summary, indent=2)};

export const EXCEL_ITOPS_MONTHLY = {json.dumps(itops_data, indent=2)};

export const EXCEL_AUTOMATION_MONTHLY = {json.dumps(auto_monthly, indent=2)};

export const EXCEL_RISK_SUMMARY = {json.dumps(risk_summary, indent=2)};

export const EXCEL_SAVINGS_DATA = {json.dumps(savings_data, indent=2)};
"""

with open('src/data/excelDataSource.ts', 'w', encoding='utf-8') as f:
    f.write(ts_content)

print("Generated src/data/excelDataSource.ts successfully!")
