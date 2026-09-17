import openpyxl
import pptx
import json
import os
from datetime import datetime

wb = openpyxl.load_workbook('DH_TechM_Managed_Services_BI_Source.xlsx', data_only=True)

def serialize_val(v):
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    return v

# 1. MOM Actions
mom_sheet = wb['02_MOM_Actions']
mom_rows = list(mom_sheet.iter_rows(values_only=True))
header_idx = -1
for i, row in enumerate(mom_rows):
    if row[0] == 'Action_No':
        header_idx = i
        break

mom_actions = []
if header_idx != -1:
    headers = mom_rows[header_idx]
    for row in mom_rows[header_idx+1:]:
        if row[0] is not None:
            item = {}
            for h, val in zip(headers, row):
                if h:
                    item[str(h)] = serialize_val(val)
            mom_actions.append(item)

print(f"Loaded {len(mom_actions)} MOM Actions")

# 2. Slide Index / Agenda
agenda_sheet = wb['01_Slide_Index']
agenda_rows = [[serialize_val(c) for c in r] for r in agenda_sheet.iter_rows(values_only=True)]

# 3. Qualys Project
qualys_sheet = wb['21_Qualys_Project']
qualys_rows = [[serialize_val(c) for c in r] for r in qualys_sheet.iter_rows(values_only=True)]

# 4. Automation Use Cases
auto_sheet = wb['12_Automation_UseCases']
auto_rows = [[serialize_val(c) for c in r] for r in auto_sheet.iter_rows(values_only=True)]

# 5. ITOps Monthly
itops_sheet = wb['03_ITOps_Monthly']
itops_rows = [[serialize_val(c) for c in r] for r in itops_sheet.iter_rows(values_only=True)]

# 6. Ticket Closure
ticket_sheet = wb['05_Ticket_Closure']
ticket_rows = [[serialize_val(c) for c in r] for r in ticket_sheet.iter_rows(values_only=True)]

# 7. Risk Dashboard & Overdue
risk_sheet = wb['22_Risk_Dashboard']
risk_rows = [[serialize_val(c) for c in r] for r in risk_sheet.iter_rows(values_only=True)]
risk_od_sheet = wb['23_Risk_Overdue']
risk_od_rows = [[serialize_val(c) for c in r] for r in risk_od_sheet.iter_rows(values_only=True)]

# 8. Value Savings
val_sheet = wb['24_Value_Savings']
val_rows = [[serialize_val(c) for c in r] for r in val_sheet.iter_rows(values_only=True)]

# 9. AIOps Roadmap
aiops_sheet = wb['25_AIOps_Roadmap']
aiops_rows = [[serialize_val(c) for c in r] for r in aiops_sheet.iter_rows(values_only=True)]

all_bi_data = {
    "mom_actions": mom_actions,
    "agenda": agenda_rows,
    "qualys": qualys_rows,
    "automation_usecases": auto_rows,
    "itops_monthly": itops_rows,
    "ticket_closure": ticket_rows,
    "risk_dashboard": risk_rows,
    "risk_overdue": risk_od_rows,
    "value_savings": val_rows,
    "aiops_roadmap": aiops_rows
}

with open('full_bi_data.json', 'w', encoding='utf-8') as f:
    json.dump(all_bi_data, f, indent=2, ensure_ascii=False)

print("Saved full_bi_data.json successfully")
